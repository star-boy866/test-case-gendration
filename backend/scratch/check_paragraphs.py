import docx

doc = docx.Document('backend/tests/fixtures/golden_sources/CR 18175 PRV-INT-027 UT DOCUMENT 1.docx')
for i, p in enumerate(doc.paragraphs):
    if 'Scenario 1:' in p.text or 'Scenario 2:' in p.text:
        print(f"P[{i}]: {p.text}")
        for j in range(i, min(i+8, len(doc.paragraphs))):
            pj = doc.paragraphs[j]
            has_drawing = any('drawing' in r._r.xml for r in pj.runs)
            print(f"   P[{j}]: text='{pj.text.strip()}' has_drawing={has_drawing}")
