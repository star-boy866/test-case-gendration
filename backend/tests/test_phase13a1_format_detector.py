"""
Phase 13A.1 Unit Tests: Auto Detect DSD Format & Safety Guardrails.

Verifies:
1. ND format detection with high confidence from ND markers.
2. NH format detection with high confidence from NH markers.
3. AK format detection with high confidence from AK markers.
4. Unknown format detection when no state markers are present.
5. POST /api/cognos/detect-dsd-format endpoint.
6. Safety guardrail in /api/cognos/upload-and-generate preventing ND/AK from silently routing through NH parser.
"""

import io
import pytest
from pathlib import Path
import docx
from fastapi.testclient import TestClient

from app.main import app
from app.core.rbac import get_current_user, CurrentUser
from app.cognos.detection.dsd_format_detector import (
    DSDFormatDetector,
    DSDProfile,
    DetectionConfidence,
    DetectionState,
)


@pytest.fixture
def mock_tester_user():
    def _override():
        return CurrentUser(username="tester", role="tester")
    app.dependency_overrides[get_current_user] = _override
    yield
    app.dependency_overrides.pop(get_current_user, None)


def _create_mock_docx(text_lines):
    doc = docx.Document()
    for line in text_lines:
        doc.add_paragraph(line)
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_detector_identifies_nd_format(tmp_path):
    """Test ND detection for North Dakota MMIS document markers."""
    doc_file = tmp_path / "TPL Rejection Error Handling Report Template (2).docx"
    doc = docx.Document()
    doc.add_paragraph("North Dakota Medicaid Systems Project MMIS Report Definition")
    doc.add_paragraph("State of North Dakota Department of Human Services")
    doc.save(str(doc_file))

    res = DSDFormatDetector.detect_format(doc_file, doc_file.name)
    assert res.detected_format == "ND"
    assert res.confidence == "High"
    assert res.display_name == "ND — North Dakota"
    assert res.status == "DETECTED"


def test_detector_identifies_nh_format(tmp_path):
    """Test NH detection for New Hampshire MMIS document markers."""
    doc_file = tmp_path / "PRV-INT-027.docx"
    doc = docx.Document()
    doc.add_paragraph("NH MMIS REPORT DEFINITION")
    doc.add_paragraph("New Hampshire Department of Health and Human Services")
    doc.save(str(doc_file))

    res = DSDFormatDetector.detect_format(doc_file, doc_file.name)
    assert res.detected_format == "NH"
    assert res.confidence == "High"
    assert res.display_name == "NH — New Hampshire"
    assert res.status == "DETECTED"


def test_detector_identifies_ak_format(tmp_path):
    """Test AK detection for Alaska Cognos DSD document markers."""
    doc_file = tmp_path / "AK-REP-001.docx"
    doc = docx.Document()
    doc.add_paragraph("Alaska Cognos DSD")
    doc.add_paragraph("State of Alaska MMIS Report")
    doc.save(str(doc_file))

    res = DSDFormatDetector.detect_format(doc_file, doc_file.name)
    assert res.detected_format == "AK"
    assert res.confidence == "High"
    assert res.display_name == "AK — Alaska"
    assert res.status == "DETECTED"


def test_detector_identifies_unknown_format(tmp_path):
    """Test Unknown format detection for generic document."""
    doc_file = tmp_path / "generic_doc.docx"
    doc = docx.Document()
    doc.add_paragraph("Standard Corporate Quarterly Performance Review")
    doc.add_paragraph("General table layout")
    doc.save(str(doc_file))

    res = DSDFormatDetector.detect_format(doc_file, doc_file.name)
    assert res.detected_format == "UNKNOWN"
    assert res.confidence == "None"
    assert res.status == "UNKNOWN"
    assert res.message is not None and "Unable to determine DSD format" in res.message


def test_detect_dsd_format_api_endpoint(mock_tester_user, tmp_path):
    """Verify POST /api/cognos/detect-dsd-format returns accurate detection payload."""
    client = TestClient(app)
    docx_buf = _create_mock_docx([
        "North Dakota Medicaid Systems Project",
        "TPL Rejection Error Handling Report",
    ])

    response = client.post(
        "/api/cognos/detect-dsd-format",
        files={"file": ("TPL Rejection Error Handling Report Template (2).docx", docx_buf.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["detected_format"] == "ND"
    assert data["confidence"] == "High"
    assert data["display_name"] == "ND — North Dakota"
    assert data["status"] == "DETECTED"


def test_upload_and_generate_unknown_safety_guardrail(mock_tester_user):
    """Verify unknown format upload fails with explicit safety guardrail message."""
    client = TestClient(app)
    docx_buf = _create_mock_docx([
        "Corporate Annual Evaluation Report",
        "Section 1: Performance",
    ])

    response = client.post(
        "/api/cognos/upload-and-generate",
        data={"dsd_profile": "AUTO"},
        files={"file": ("unknown_doc.docx", docx_buf.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    assert "Unable to determine DSD format. Please select NH / ND / AK manually." in response.json()["detail"]


def test_upload_and_generate_ak_safety_guardrail(mock_tester_user):
    """Verify AK upload fails with explicit safety guardrail message."""
    client = TestClient(app)
    docx_buf = _create_mock_docx([
        "Alaska Cognos DSD",
        "Alaska Department of Health",
    ])

    response = client.post(
        "/api/cognos/upload-and-generate",
        data={"dsd_profile": "AK"},
        files={"file": ("AK-REP-001.docx", docx_buf.read(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    assert "AK format detected. AK processing is not enabled yet." in response.json()["detail"]
